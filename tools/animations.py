"""PowerPoint click-build extraction and per-step slide state rendering.

PNG conversion of a slide loses PowerPoint's build order. This module reads the
slide's ``p:timing`` main sequence (the click-by-click animation list), works out
which pieces of the slide are visible after every click, and lets the bundle
builder render one image per click state.

What is preserved: the ORDER of the builds, the visible STATE after every click,
paragraph builds inside text boxes, and chart builds by series, category or
point. What is NOT reproduced: the motion itself (fly, wipe, spin ...) is shown as
a calm cross-fade between states; emphasis effects, motion paths, SmartArt builds
and interactive triggers are reported, never silently pretended.
"""
from __future__ import annotations

import copy
import re
from dataclasses import dataclass, field

from lxml import etree

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
NS = {"p": P, "a": A, "c": C}


def q(prefix: str, name: str) -> str:
    return f"{{{NS[prefix]}}}{name}"


PRESET_NAMES = {
    1: "appear", 2: "fly", 3: "blinds", 4: "box", 5: "checkerboard", 6: "circle",
    7: "crawl", 8: "diamond", 9: "dissolve", 10: "fade", 12: "peek", 13: "plus",
    14: "random-bars", 16: "split", 17: "stretch", 18: "strips", 21: "wheel",
    22: "wipe", 23: "zoom", 26: "bounce", 53: "grow-turn", 55: "zoom-in",
}

# Features whose presence means the click list cannot be trusted at all.
BLOCKING = {"interactive", "smartart", "ole-chart"}


@dataclass
class Effect:
    action: str            # "show" | "hide"
    units: list            # atomic units touched, see unit_* helpers
    start_ms: int = 0
    dur_ms: int = 0
    preset: str = "appear"
    target: str = "shape"  # shape | paragraph | chart


@dataclass
class Click:
    auto_ms: int | None = None  # None = waits for a click; else starts by itself
    effects: list = field(default_factory=list)

    @property
    def duration_ms(self) -> int:
        return max((e.start_ms + e.dur_ms for e in self.effects), default=0)


@dataclass
class BuildModel:
    clicks: list
    initially_hidden: set
    unsupported: list
    blocking: list

    @property
    def usable(self) -> bool:
        return bool(self.clicks) and not self.blocking

    def visible_after(self, step: int) -> dict:
        """Map of unit -> visible after `step` clicks (0 = slide start)."""
        state = {u: False for u in self.initially_hidden}
        for click in self.clicks[:step]:
            for effect in click.effects:
                for unit in effect.units:
                    state[unit] = effect.action == "show"
        return state


def _int(value, default=0) -> int:
    try:
        return int(str(value))
    except (TypeError, ValueError):
        return default


def _cond_delay(ctn) -> int | None:
    """Delay of a cTn start condition in ms; None for 'indefinite' (waits for a click)."""
    cond = ctn.find("p:stCondLst/p:cond", NS)
    if cond is None:
        return 0
    delay = cond.get("delay")
    if delay == "indefinite":
        return None
    return _int(delay)


def _shape_elements(slide) -> dict:
    found = {}
    for cnv in slide._element.iter(q("p", "cNvPr")):
        ident = _int(cnv.get("id"), -1)
        parent = cnv.getparent()
        if ident >= 0 and parent is not None and parent.getparent() is not None:
            found.setdefault(ident, parent.getparent())
    return found


def _chart_shape(slide, spid):
    from pptx.shapes.graphfrm import GraphicFrame
    for shape in _walk(slide.shapes):
        if shape.shape_id == spid and isinstance(shape, GraphicFrame) and shape.has_chart:
            return shape
    return None


def _walk(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == 6:
            yield from _walk(shape.shapes)


def _chart_dimensions(slide, spid):
    shape = _chart_shape(slide, spid)
    if shape is None:
        return None
    space = shape.chart._chartSpace
    series = space.findall(".//c:ser", NS)
    if not series:
        return None
    counts = []
    for ser in series:
        pt_count = ser.find("c:val/c:numRef/c:numCache/c:ptCount", NS)
        if pt_count is None:
            pt_count = ser.find("c:cat//c:ptCount", NS)
        counts.append(_int(pt_count.get("val")) if pt_count is not None else 0)
    kinds = {el.tag.split("}")[1] for el in space.find(".//c:plotArea", NS)
             if el.tag.endswith("Chart")}
    if kinds & {"scatterChart", "bubbleChart", "stockChart", "surfaceChart", "surface3DChart"}:
        return None
    return len(series), max(counts) if counts else 0


def _targets(effect_ctn):
    """Yield (kind, spid, extra) for every distinct target of an effect node."""
    seen = set()
    for tgt in effect_ctn.iter(q("p", "spTgt")):
        spid = _int(tgt.get("spid"), -1)
        key = None
        extra = None
        if tgt.find("p:txEl", NS) is not None:
            rng = tgt.find("p:txEl/p:pRg", NS)
            if rng is not None:
                key = ("paragraph", spid, _int(rng.get("st")), _int(rng.get("end")))
            else:
                key = ("unsupported-text", spid, 0, 0)
        elif tgt.find("p:graphicEl", NS) is not None:
            chart = tgt.find("p:graphicEl/a:chart", NS)
            if chart is not None:
                key = ("chart", spid, chart.get("bldStep", "allAtOnce"),
                       _int(chart.get("seriesIdx"), -3), _int(chart.get("categoryIdx"), -3))
            else:
                key = ("smartart", spid, 0, 0)
        elif tgt.find("p:oleChartEl", NS) is not None:
            key = ("ole-chart", spid, 0, 0)
        elif tgt.find("p:subSp", NS) is not None:
            key = ("unsupported-sub", spid, 0, 0)
        else:
            key = ("shape", spid)
        if key not in seen:
            seen.add(key)
            yield key


def _filter_name(ctn, preset_id) -> str:
    effect = next(ctn.iter(q("p", "animEffect")), None)
    if effect is not None and effect.get("filter"):
        return re.sub(r"\(.*\)", "", effect.get("filter")).strip()
    return PRESET_NAMES.get(preset_id, "effect")


def _duration(ctn) -> int:
    best = 0
    for node in ctn.iter(q("p", "cTn")):
        dur = node.get("dur")
        if dur and dur.isdigit():
            best = max(best, int(dur))
    return best


def parse_build_model(slide) -> BuildModel | None:
    """Return the click-by-click build model, or None if the slide has no builds."""
    timing = slide._element.find("p:timing", NS)
    if timing is None:
        return None
    unsupported: list[str] = []
    blocking: list[str] = []
    if any(True for _ in timing.iter(q("p", "seq"))
           if _.find("p:cTn", NS) is not None and _.find("p:cTn", NS).get("nodeType") == "interactiveSeq"):
        blocking.append("interactive")
    main = None
    for ctn in timing.iter(q("p", "cTn")):
        if ctn.get("nodeType") == "mainSeq":
            main = ctn
            break
    if main is None:
        return BuildModel([], set(), unsupported, blocking) if blocking else None
    child_list = main.find("p:childTnLst", NS)
    click_nodes = child_list.findall("p:par", NS) if child_list is not None else []
    charts_dims: dict[int, tuple | None] = {}
    clicks: list[Click] = []
    first_action: dict = {}
    order: list = []

    def expand(key):
        kind = key[0]
        if kind == "shape":
            return "shape", [("shape", key[1])]
        if kind == "paragraph":
            _, spid, start, end = key
            return "paragraph", [("para", spid, i) for i in range(start, max(start, end) + 1)]
        if kind == "chart":
            _, spid, step, series, category = key
            if spid not in charts_dims:
                charts_dims[spid] = _chart_dimensions(slide, spid)
            dims = charts_dims[spid]
            if dims is None:
                unsupported.append("chart-type")
                return "chart", [("shape", spid)]
            n_series, n_cat = dims
            if step == "gridLegend":
                return "chart", [("chartbg", spid)]
            if step == "series":
                return "chart", [("pt", spid, s, c) for s in ([series] if series >= 0 else [])
                                 for c in range(n_cat)]
            if step == "category":
                return "chart", [("pt", spid, s, c) for c in ([category] if category >= 0 else [])
                                 for s in range(n_series)]
            if step in ("ptInSeries", "ptInCategory") and series >= 0 and category >= 0:
                return "chart", [("pt", spid, series, category)]
            if step == "allAtOnce":
                return "chart", [("shape", spid)]
            unsupported.append("chart-build")
            return "chart", [("shape", spid)]
        blocking.append("smartart" if kind == "smartart" else
                        "ole-chart" if kind == "ole-chart" else "unsupported-target")
        return "shape", []

    for index, click_node in enumerate(click_nodes):
        ctn = click_node.find("p:cTn", NS)
        if ctn is None:
            continue
        waits = _cond_delay(ctn)
        auto = None if waits is None else waits
        if ctn.find("p:stCondLst/p:cond[@evt='onBegin']", NS) is not None:
            auto = 0
        click = Click(auto_ms=auto)
        for group in ctn.findall("p:childTnLst/p:par", NS):
            gctn = group.find("p:cTn", NS)
            if gctn is None:
                continue
            group_delay = _cond_delay(gctn) or 0
            for node in gctn.findall("p:childTnLst/p:par", NS):
                ectn = node.find("p:cTn", NS)
                if ectn is None or not ectn.get("presetClass"):
                    continue
                klass = ectn.get("presetClass")
                start = group_delay + (_cond_delay(ectn) or 0)
                preset_id = _int(ectn.get("presetID"), 0)
                if klass == "mediacall":
                    continue  # Videos are played by the player itself.
                if klass in ("emph", "path", "verb"):
                    unsupported.append("emphasis-or-path-motion")
                    continue
                action = "show" if klass == "entr" else "hide"
                # Some custom effects set visibility explicitly; trust that when present.
                for setter in ectn.iter(q("p", "set")):
                    attr = setter.find("p:cBhvr/p:attrNameLst/p:attrName", NS)
                    value = setter.find("p:to/p:strVal", NS)
                    if attr is not None and attr.text == "style.visibility" and value is not None:
                        action = "show" if value.get("val") == "visible" else "hide"
                        break
                for key in _targets(ectn):
                    target, units = expand(key)
                    if not units:
                        continue
                    for unit in units:
                        if unit not in first_action:
                            first_action[unit] = action
                            order.append(unit)
                    click.effects.append(Effect(
                        action=action, units=units, start_ms=start,
                        dur_ms=_duration(ectn), preset=_filter_name(ectn, preset_id),
                        target=target))
        if click.effects:
            clicks.append(click)
    # `chartbg` hidden at the start means the whole frame is not drawn yet.
    hidden = {u for u in order if first_action[u] == "show"}
    for chart_unit in [u for u in hidden if u[0] == "pt"]:
        pass
    return BuildModel(clicks, hidden, sorted(set(unsupported)), sorted(set(blocking)))


# ---------------------------------------------------------------------------
# Rendering a state: modify a COPY of the slide so LibreOffice draws only what
# PowerPoint would show after N clicks.
# ---------------------------------------------------------------------------

_BU = {"buClrTx", "buClr", "buSzTx", "buSzPct", "buSzPts", "buFontTx", "buFont",
       "buNone", "buAutoNum", "buChar", "buBlip"}
_RPR_FILLS = {"noFill", "solidFill", "gradFill", "blipFill", "pattFill", "grpFill"}


def _hide_paragraph(p_el):
    ppr = p_el.find("a:pPr", NS)
    if ppr is None:
        ppr = etree.SubElement(p_el, q("a", "pPr"))
        p_el.insert(0, ppr)
    for child in list(ppr):
        if etree.QName(child).localname in _BU:
            ppr.remove(child)
    marker = etree.Element(q("a", "buNone"))
    for index, child in enumerate(ppr):
        if etree.QName(child).localname in ("tabLst", "defRPr", "extLst"):
            ppr.insert(index, marker)
            break
    else:
        ppr.append(marker)
    for run in p_el:
        name = etree.QName(run).localname
        if name not in ("r", "fld", "br"):
            continue
        rpr = run.find("a:rPr", NS)
        if rpr is None:
            rpr = etree.Element(q("a", "rPr"))
            run.insert(0, rpr)
        for child in list(rpr):
            if etree.QName(child).localname in _RPR_FILLS:
                rpr.remove(child)
        position = 1 if rpr.find("a:ln", NS) is not None else 0
        rpr.insert(position, etree.Element(q("a", "noFill")))


_AFTER_DPT = {"dLbls", "trendline", "errBars", "cat", "val", "xVal", "yVal",
              "bubbleSize", "smooth", "shape", "extLst"}


def _hide_chart_point(ser, idx):
    parent = etree.QName(ser.getparent()).localname
    dpt = None
    for existing in ser.findall("c:dPt", NS):
        if _int(existing.find("c:idx", NS).get("val"), -1) == idx:
            dpt = existing
            break
    if dpt is None:
        dpt = etree.Element(q("c", "dPt"))
        etree.SubElement(dpt, q("c", "idx")).set("val", str(idx))
        for index, child in enumerate(ser):
            if etree.QName(child).localname in _AFTER_DPT:
                ser.insert(index, dpt)
                break
        else:
            ser.append(dpt)
    for child in list(dpt):
        if etree.QName(child).localname in ("spPr", "marker"):
            dpt.remove(child)
    if parent in ("lineChart", "radarChart"):
        marker = etree.SubElement(dpt, q("c", "marker"))
        etree.SubElement(marker, q("c", "symbol")).set("val", "none")
    sppr = etree.SubElement(dpt, q("c", "spPr"))
    etree.SubElement(sppr, q("a", "noFill"))
    line = etree.SubElement(sppr, q("a", "ln"))
    etree.SubElement(line, q("a", "noFill"))
    labels = ser.find("c:dLbls", NS)
    if labels is not None:
        for old in labels.findall("c:dLbl", NS):
            if _int(old.find("c:idx", NS).get("val"), -1) == idx:
                labels.remove(old)
        entry = etree.Element(q("c", "dLbl"))
        etree.SubElement(entry, q("c", "idx")).set("val", str(idx))
        etree.SubElement(entry, q("c", "delete")).set("val", "1")
        labels.insert(0, entry)


def apply_state(slide, model: BuildModel, step: int):
    """Mutate `slide` (a throwaway copy) to show exactly what is visible after `step` clicks."""
    visible = model.visible_after(step)
    elements = _shape_elements(slide)
    removed = set()
    for unit, shown in visible.items():
        if shown:
            continue
        kind = unit[0]
        if kind in ("shape", "chartbg"):
            element = elements.get(unit[1])
            if element is not None and element.getparent() is not None:
                element.getparent().remove(element)
                removed.add(unit[1])
    for unit, shown in visible.items():
        if shown or unit[1] in removed:
            continue
        kind = unit[0]
        element = elements.get(unit[1])
        if element is None:
            continue
        if kind == "para":
            paragraphs = element.findall(".//a:p", NS)
            if unit[2] < len(paragraphs):
                _hide_paragraph(paragraphs[unit[2]])
        elif kind == "pt":
            shape = _chart_shape(slide, unit[1])
            if shape is None:
                continue
            sers = shape.chart._chartSpace.findall(".//c:ser", NS)
            if unit[2] < len(sers):
                _hide_chart_point(sers[unit[2]], unit[3])


def describe_steps(model: BuildModel, image_for) -> list:
    """Manifest entries; entry 0 is the slide at rest, entry k the state after click k."""
    steps = [{"image": image_for(0), "auto": False, "delayMs": 0, "durationMs": 0, "effects": []}]
    for number, click in enumerate(model.clicks, start=1):
        steps.append({
            "image": image_for(number),
            "auto": click.auto_ms is not None,
            "delayMs": click.auto_ms or 0,
            "durationMs": min(click.duration_ms, 5000),
            "effects": [{"action": e.action, "target": e.target, "effect": e.preset,
                         "startMs": min(e.start_ms, 60000), "durMs": min(e.dur_ms, 5000)}
                        for e in click.effects[:40]],
        })
    return steps


def render_steps(pptx, out_dir, width, models: dict) -> dict:
    """Render slides/steps images: one LibreOffice pass per click number, animated slides only.

    `models` maps slide number -> usable BuildModel. Returns {slide: [relative image paths]}.
    """
    import shutil
    import subprocess
    from pathlib import Path
    from pptx import Presentation

    out_dir = Path(out_dir)
    steps_dir = out_dir / "steps"
    shutil.rmtree(steps_dir, ignore_errors=True)
    steps_dir.mkdir(parents=True)
    images: dict[int, list] = {number: [] for number in models}
    most = max(len(model.clicks) for model in models.values())
    keep = sorted(models)
    work = out_dir / "_steps_render"
    try:
        for step in range(most + 1):
            shutil.rmtree(work, ignore_errors=True)
            work.mkdir(parents=True)
            deck = Presentation(str(pptx))
            ids = deck.slides._sldIdLst
            rendered = []
            for number, slide_id in enumerate(list(ids), start=1):
                if number not in models:
                    deck.part.drop_rel(slide_id.rId)
                    ids.remove(slide_id)
                    continue
                model = models[number]
                if step <= len(model.clicks):
                    apply_state(_slide_by_id(deck, slide_id), model, step)
                    rendered.append(number)
            deck.save(str(work / "steps.pptx"))
            subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(work),
                            str(work / "steps.pptx")], check=True, capture_output=True, timeout=600)
            subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(width), "-scale-to-y", "-1",
                            str(work / "steps.pdf"), str(work / "page")], check=True, timeout=600)
            pages = sorted(work.glob("page-*.png"))
            if len(pages) != len(keep):
                raise RuntimeError(f"expected {len(keep)} pages for build step {step}, got {len(pages)}")
            for number, page in zip(keep, pages):
                if number in rendered:
                    name = f"steps/{number:03d}-{step:02d}.png"
                    shutil.move(str(page), out_dir / name)
                    images[number].append(name)
    finally:
        shutil.rmtree(work, ignore_errors=True)
    return images


def _slide_by_id(deck, slide_id):
    return deck.part.related_part(slide_id.rId).slide
