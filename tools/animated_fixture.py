"""Generate a synthetic PPTX whose timing XML follows PowerPoint's click-build structure.

Used only by tests: slide 1 builds three bullet paragraphs (click, click, with-previous),
slide 2 builds a clustered column chart by series, slide 3 has no animation.
"""
import sys
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

P = "http://schemas.openxmlformats.org/presentationml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"


def effect(ident, spid, node_type, preset=10, klass="entr", target_xml=None, delay=0, dur=500):
    target = target_xml or f'<p:spTgt spid="{spid}"/>'
    value = "visible" if klass == "entr" else "hidden"
    return (
        f'<p:par><p:cTn id="{ident}" presetID="{preset}" presetClass="{klass}" presetSubtype="0" '
        f'fill="hold" nodeType="{node_type}"><p:stCondLst><p:cond delay="{delay}"/></p:stCondLst>'
        f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{ident + 1}" dur="1" fill="hold">'
        f'<p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn><p:tgtEl>{target}</p:tgtEl>'
        f'<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr>'
        f'<p:to><p:strVal val="{value}"/></p:to></p:set>'
        f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{ident + 2}" dur="{dur}"/>'
        f'<p:tgtEl>{target}</p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par>'
    )


def click(ident, effects_xml, auto=False):
    cond = '<p:cond delay="indefinite"/>' + ('<p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond>' if auto else "")
    return (
        f'<p:par><p:cTn id="{ident}" fill="hold"><p:stCondLst>{cond}</p:stCondLst><p:childTnLst>'
        f'<p:par><p:cTn id="{ident + 1}" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
        f'<p:childTnLst>{effects_xml}</p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par>'
    )


def add_timing(slide, clicks_xml):
    xml = (
        f'<p:timing xmlns:p="{P}" xmlns:a="{A}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" '
        f'restart="never" nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek">'
        f'<p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>{clicks_xml}</p:childTnLst>'
        f'</p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
        f'</p:prevCondLst><p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl>'
        f'</p:cond></p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>'
    )
    slide._element.append(etree.fromstring(xml))


def build(path):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(11), Inches(4))
    frame = box.text_frame
    frame.text = "النقطة الأولى"
    for text in ("النقطة الثانية", "النقطة الثالثة"):
        frame.add_paragraph().text = text
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(40)
    sid = box.shape_id
    rng = lambda i: f'<p:spTgt spid="{sid}"><p:txEl><p:pRg st="{i}" end="{i}"/></p:txEl></p:spTgt>'
    add_timing(slide,
        click(3, effect(5, sid, "clickEffect", target_xml=rng(0))) +
        click(10, effect(12, sid, "clickEffect", target_xml=rng(1)) +
              effect(15, sid, "withEffect", target_xml=rng(2), delay=250)))

    slide2 = deck.slides.add_slide(deck.slide_layouts[6])
    data = CategoryChartData()
    data.categories = ["أ", "ب", "ج"]
    data.add_series("الأولى", (40, 30, 20))
    data.add_series("الثانية", (10, 50, 35))
    gf = slide2.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1),
                                 Inches(11), Inches(5.5), data)
    gid = gf.shape_id
    chart_t = lambda step, s=-3, c=-3: (
        f'<p:spTgt spid="{gid}"><p:graphicEl><a:chart seriesIdx="{s}" categoryIdx="{c}" '
        f'bldStep="{step}"/></p:graphicEl></p:spTgt>')
    add_timing(slide2,
        click(3, effect(5, gid, "clickEffect", target_xml=chart_t("series", 0))) +
        click(10, effect(12, gid, "clickEffect", target_xml=chart_t("series", 1))))

    slide3 = deck.slides.add_slide(deck.slide_layouts[6])
    slide3.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text_frame.text = "بلا حركة"
    deck.save(path)
    return path


if __name__ == "__main__":
    print(build(Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/alfa-animated.pptx")))
