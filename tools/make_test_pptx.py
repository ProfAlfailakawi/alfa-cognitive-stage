"""Simple public-domain test lecture generated only inside CI."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pathlib import Path
import sys
out = Path(sys.argv[1] if len(sys.argv) > 1 else "/tmp/alfa-ci-lecture.pptx")
deck = Presentation()
deck.slide_width = Inches(13.333)
deck.slide_height = Inches(7.5)
for i, phrase in enumerate(["الألعاب التعليمية", "التصويت والمشاركة"], start=1):
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(10), Inches(2.3))
    paragraph = shape.text_frame.paragraphs[0]
    paragraph.text = phrase
    paragraph.font.size = Pt(36)
    slide.notes_slide.notes_text_frame.text = "ملاحظات تجربة " + str(i)
deck.save(out)
print(out)
