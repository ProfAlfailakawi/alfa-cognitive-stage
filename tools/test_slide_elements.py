"""Conservative PowerPoint element geometry for X-Ray, never infer group/rotation."""
import unittest
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.shapes import MSO_SHAPE

import build_bundle


class XrayBoundsTests(unittest.TestCase):
    def setUp(self):
        self.deck = Presentation()
        self.slide = self.deck.slides.add_slide(self.deck.slide_layouts[6])

    def test_rtl_text_coordinates_are_stable(self):
        shape = self.slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        shape.text = "الدافعية والتعلّم النشط"
        items = build_bundle.slide_elements(self.slide, self.deck.slide_width, self.deck.slide_height, 1)
        match = [x for x in items if x["text"] == shape.text]
        self.assertEqual(len(match), 1)
        self.assertAlmostEqual(match[0]["x"], float(Inches(1) / self.deck.slide_width), places=4)
        self.assertAlmostEqual(match[0]["w"], float(Inches(4) / self.deck.slide_width), places=4)
        self.assertTrue(match[0]["id"].startswith("s1-e"))

    def test_four_by_three_slide_uses_original_dimensions(self):
        self.deck.slide_width, self.deck.slide_height = Inches(10), Inches(7.5)
        box = self.slide.shapes.add_textbox(Inches(5), Inches(2), Inches(2), Inches(1))
        box.text = "English and العربية 2026"
        items = build_bundle.slide_elements(self.slide, self.deck.slide_width, self.deck.slide_height, 2)
        row = next(item for item in items if item["text"] == box.text)
        self.assertAlmostEqual(row["x"], .5, places=5)
        self.assertAlmostEqual(row["y"], 2/7.5, places=5)
        self.assertAlmostEqual(row["w"], .2, places=5)

    def test_rotated_shape_is_not_focusable(self):
        shape = self.slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1),
            Inches(2), Inches(1))
        shape.rotation = 42
        items = build_bundle.slide_elements(self.slide, self.deck.slide_width, self.deck.slide_height, 1)
        self.assertFalse(any(row["id"] == f"s1-e{shape.shape_id}" for row in items))

    def test_off_slide_object_is_skipped(self):
        shape = self.slide.shapes.add_textbox(Inches(13), Inches(1), Inches(2), Inches(1))
        shape.text = "خارج الشريحة"
        items = build_bundle.slide_elements(self.slide, self.deck.slide_width, self.deck.slide_height, 2)
        self.assertFalse(any(row["text"] == shape.text for row in items))
        self.assertLessEqual(len(items), 120)


if __name__ == "__main__":
    unittest.main()
