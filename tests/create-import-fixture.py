"""Creates a small real PPTX including an embedded MP4 for CI import checks."""
from pathlib import Path
import sys
from pptx import Presentation
from pptx.util import Inches
from PIL import Image

dest = Path(sys.argv[1])
mp4 = Path(sys.argv[2])
poster = dest.with_suffix(".png")
Image.new("RGB", (320, 240), "#27465f").save(poster)
ppt = Presentation()
ppt.slide_width = Inches(13.333)
ppt.slide_height = Inches(7.5)
slide1 = ppt.slides.add_slide(ppt.slide_layouts[6])
slide1.shapes.add_textbox(Inches(1), Inches(1), Inches(9), Inches(1.3)).text = "الألعاب التعليمية"
slide1.notes_slide.notes_text_frame.text = "شاهدوا الفيديو في الشريحة الثانية"
slide2 = ppt.slides.add_slide(ppt.slide_layouts[6])
slide2.shapes.add_textbox(Inches(1), Inches(.4), Inches(9), Inches(1)).text = "مقطع فيديو تعليمي"
slide2.shapes.add_movie(
    str(mp4), Inches(1), Inches(1.6), Inches(8), Inches(4.5),
    poster_frame_image=str(poster), mime_type="video/mp4"
)
ppt.save(dest)
print(dest)
